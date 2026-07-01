"""
MD → DOCX 转换脚本 v2.1
按照滇池学院期末考试论文格式规范生成 .docx 文件

格式要点：
- 摘要标题：黑体 小四 加粗 居中, 段前/段后1行, 单倍行距
- 摘要正文：宋体/TNR 五号, 两端对齐, 首行缩进2字符, 1.5倍行距
- 关键词：黑体 小四 加粗 "关键词：" + 宋体/TNR 五号 内容
- 章标题(第X章)：黑体 四号 加粗 居中, 每章另起新页
- 节标题(X.X)：黑体 小四 加粗 左对齐
- 小节标题(X.X.X)：宋体 小四 加粗 左对齐
- 正文：宋体/TNR 小四, 两端对齐, 首行缩进2字符, 1.5倍行距
- 图注：黑体/TNR 五号 居中
- 表注：黑体/TNR 五号 居中, 三线表
- 表内容：宋体/TNR 五号, 1.25倍行距, 居中
- 参考文献标题：黑体 四号 加粗 居中, 另起新页
- 参考文献条目：宋体/TNR 小四, 悬挂缩进2字符, 1.5倍行距
- 页码：TNR 五号, 页面底部居中, 阿拉伯数字从1开始
"""
import re
import os
import sys
sys.stdout.reconfigure(encoding='utf-8')

from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE_DIR = r"E:\work\Claude code default\自然语言处理期末"
MD_FILE = os.path.join(BASE_DIR, "论文", "学术论文.md")
DOCX_FILE = os.path.join(BASE_DIR, "论文", "学术论文.docx")

# ======================== 字号映射 ========================
# 三号=16pt, 四号=14pt, 小四=12pt, 五号=10.5pt
SZ_SANHAO = Pt(16)
SZ_SIHAO  = Pt(14)
SZ_XIAOSI = Pt(12)
SZ_WUHAO  = Pt(10.5)

doc = Document()

# ======================== 页面设置 (A4) ========================
section = doc.sections[0]
section.page_width  = Cm(21.0)
section.page_height = Cm(29.7)
section.top_margin    = Cm(2.54)
section.bottom_margin = Cm(2.54)
section.left_margin   = Cm(3.18)
section.right_margin  = Cm(3.18)

# ======================== 基础样式 ========================
style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = SZ_XIAOSI
style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
style.paragraph_format.line_spacing = 1.5

# ======================== 字体设置工具 ========================

def set_font(run, cn='宋体', en='Times New Roman', size=SZ_XIAOSI, bold=False):
    """统一设置中英文字体"""
    run.font.name = en
    run.font.size = size
    run.bold = bold
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.insert(0, rFonts)
    rFonts.set(qn('w:eastAsia'), cn)
    rFonts.set(qn('w:ascii'), en)
    rFonts.set(qn('w:hAnsi'), en)

def add_keep_next(paragraph):
    """设置'与下段同页'"""
    pPr = paragraph._p.get_or_add_pPr()
    kn = OxmlElement('w:keepNext')
    pPr.append(kn)

# ======================== 段落构建函数 ========================

def add_abstract_heading(text="摘要"):
    """摘要标题：黑体 小四 加粗 居中, 段前/后1行, 单倍行距, 大纲1级"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.line_spacing = 1.0
    pf.space_before = Pt(14)
    pf.space_after = Pt(14)
    pf.first_line_indent = Cm(0)
    run = p.add_run(text)
    set_font(run, cn='黑体', size=SZ_XIAOSI, bold=True)
    return p

def add_abstract_body(text):
    """摘要正文：宋体/TNR 五号, 两端对齐, 首行缩进2字符, 1.5倍行距"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.first_line_indent = SZ_WUHAO * 2  # 2个五号字符
    parts = re.split(r'(\*\*.*?\*\*|\*.*?\*)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = p.add_run(part[2:-2])
            set_font(run, size=SZ_WUHAO, bold=True)
        elif part.startswith('*') and part.endswith('*') and not part.startswith('**'):
            run = p.add_run(part[1:-1])
            set_font(run, size=SZ_WUHAO)
            run.italic = True
        else:
            run = p.add_run(part)
            set_font(run, size=SZ_WUHAO)
    return p

def add_keywords_line(keywords_text):
    """关键词：黑体小四'关键词：'+ 宋体五号内容, 左对齐, 段后1行"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    pf.space_before = Pt(0)
    pf.space_after = Pt(14)
    pf.first_line_indent = Cm(0)
    run_label = p.add_run('关键词：')
    set_font(run_label, cn='黑体', size=SZ_XIAOSI, bold=True)
    run_content = p.add_run(keywords_text)
    set_font(run_content, cn='宋体', size=SZ_WUHAO)
    return p

def add_first_chapter_break():
    """在摘要/关键词之后、第1章之前添加分节符（用于页码控制）"""
    new_section = doc.add_section()
    new_section.page_width  = Cm(21.0)
    new_section.page_height = Cm(29.7)
    new_section.top_margin    = Cm(2.54)
    new_section.bottom_margin = Cm(2.54)
    new_section.left_margin   = Cm(3.18)
    new_section.right_margin  = Cm(3.18)

def add_chapter_heading(text, is_first=False):
    """章标题(第X章)：黑体 四号 加粗 居中, 段前1行 段后0.5行, 1.5倍, 每章新页"""
    if is_first:
        add_first_chapter_break()
    else:
        add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    pf.space_before = Pt(14)
    pf.space_after = Pt(7)
    pf.first_line_indent = Cm(0)
    run = p.add_run(text)
    set_font(run, cn='黑体', size=SZ_SIHAO, bold=True)
    add_keep_next(p)
    return p

def add_section_heading(text):
    """节标题(X.X)：黑体 小四 加粗 左对齐, 段前0.5行 段后0, 1.5倍"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    pf.space_before = Pt(7)
    pf.space_after = Pt(0)
    pf.first_line_indent = Cm(0)
    run = p.add_run(text)
    set_font(run, cn='黑体', size=SZ_XIAOSI, bold=True)
    add_keep_next(p)
    return p

def add_subsection_heading(text):
    """小节标题(X.X.X)：宋体 小四 加粗 左对齐, 段前0.5行 段后0, 1.5倍"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    pf.space_before = Pt(7)
    pf.space_after = Pt(0)
    pf.first_line_indent = Cm(0)
    run = p.add_run(text)
    set_font(run, cn='宋体', size=SZ_XIAOSI, bold=True)
    add_keep_next(p)
    return p

def add_subsubsection_heading(text):
    """四级标题(X.X.X.X)：宋体 小四 常规 左对齐, 段前/后0, 1.5倍"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.first_line_indent = Cm(0)
    run = p.add_run(text)
    set_font(run, cn='宋体', size=SZ_XIAOSI, bold=False)
    add_keep_next(p)
    return p

def add_body_paragraph(text):
    """正文：宋体/TNR 小四, 两端对齐, 首行缩进2字符, 1.5倍, 段前/后0"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.first_line_indent = SZ_XIAOSI * 2
    parts = re.split(r'(\*\*.*?\*\*|\*.*?\*)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = p.add_run(part[2:-2])
            set_font(run, bold=True)
        elif part.startswith('*') and part.endswith('*') and not part.startswith('**'):
            run = p.add_run(part[1:-1])
            set_font(run)
            run.italic = True
        else:
            run = p.add_run(part)
            set_font(run)
    return p

def add_ref_heading(text="参考文献"):
    """参考文献标题：新页, 黑体 四号 加粗 居中, 段前1行 段后0.5行, 1.5倍"""
    add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    pf.space_before = Pt(14)
    pf.space_after = Pt(7)
    pf.first_line_indent = Cm(0)
    run = p.add_run(text)
    set_font(run, cn='黑体', size=SZ_SIHAO, bold=True)
    return p

def add_reference_entry(text):
    """参考文献条目：宋体/TNR 小四, 悬挂缩进2字符, 1.5倍, 两端对齐"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    # 悬挂缩进2字符: 左缩进2字符 + 首行缩进-2字符
    indent_2char = Cm(0.85)  # 小四(12pt) × 2 ≈ 24pt ≈ 0.85cm
    pf.left_indent = indent_2char
    pf.first_line_indent = -indent_2char
    run = p.add_run(text)
    set_font(run, cn='宋体', size=SZ_XIAOSI)
    return p

def add_page_break():
    """分页符"""
    p = doc.add_paragraph()
    run = p.add_run()
    br = OxmlElement('w:br')
    br.set(qn('w:type'), 'page')
    run._r.append(br)
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)

# ======================== 图表函数 ========================

def add_figure(image_path, caption_text):
    """图片+图注：图片居中, 图注黑体/TNR 五号 居中"""
    abs_path = os.path.normpath(os.path.join(os.path.dirname(MD_FILE), image_path))
    if os.path.exists(abs_path):
        try:
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.first_line_indent = Cm(0)
            p_img.paragraph_format.space_before = Pt(6)
            p_img.paragraph_format.space_after = Pt(3)
            run = p_img.add_run()
            run.add_picture(abs_path, width=Inches(5.0))

            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf = p_cap.paragraph_format
            pf.line_spacing = 1.5
            pf.first_line_indent = Cm(0)
            pf.space_before = Pt(0)
            pf.space_after = Pt(6)
            run = p_cap.add_run(caption_text)
            set_font(run, cn='黑体', size=SZ_WUHAO)

            add_keep_next(p_img)
        except Exception as e:
            doc.add_paragraph(f"[图片加载失败: {caption_text}] ({e})")
    else:
        doc.add_paragraph(f"[图片文件未找到: {caption_text}]")

def add_table_caption(caption_text):
    """表注：黑体/TNR 五号 居中, 与下表同页"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    pf.space_before = Pt(6)
    pf.space_after = Pt(3)
    pf.first_line_indent = Cm(0)
    run = p.add_run(caption_text)
    set_font(run, cn='黑体', size=SZ_WUHAO)
    add_keep_next(p)
    return p

def create_three_line_table(rows_data, num_cols):
    """三线表：顶线1.5pt + 表头底线0.75pt + 底线1.5pt, 无竖线"""
    table = doc.add_table(rows=len(rows_data), cols=num_cols)
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table.autofit = True

    # 填充内容
    for i, row_data in enumerate(rows_data):
        for j, cell_text in enumerate(row_data):
            if j < num_cols:
                cell = table.rows[i].cells[j]
                cell.paragraphs[0].clear()
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                pf = p.paragraph_format
                pf.line_spacing = 1.25
                pf.space_before = Pt(1)
                pf.space_after = Pt(1)
                pf.first_line_indent = Cm(0)
                run = p.add_run(cell_text)
                if i == 0:
                    set_font(run, cn='黑体', size=SZ_WUHAO, bold=True)
                else:
                    set_font(run, cn='宋体', size=SZ_WUHAO, bold=False)

    # === 三线表边框 (XML级别) ===
    tbl = table._tbl
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement('w:tblPr')
        tbl.insert(0, tblPr)

    # 移除旧边框
    for old in tblPr.findall(qn('w:tblBorders')):
        tblPr.remove(old)

    borders = OxmlElement('w:tblBorders')

    # 顶线 1.5pt (sz=12 即 12/8 = 1.5pt)
    top = OxmlElement('w:top')
    top.set(qn('w:val'), 'single')
    top.set(qn('w:sz'), '12')
    top.set(qn('w:space'), '0')
    top.set(qn('w:color'), '000000')
    borders.append(top)

    # 底线 1.5pt
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '12')
    bottom.set(qn('w:space'), '0')
    bottom.set(qn('w:color'), '000000')
    borders.append(bottom)

    # 左右竖线: 无
    for side in ['left', 'right']:
        elem = OxmlElement(f'w:{side}')
        elem.set(qn('w:val'), 'none')
        elem.set(qn('w:sz'), '0')
        elem.set(qn('w:space'), '0')
        elem.set(qn('w:color'), 'auto')
        borders.append(elem)

    # 内部横线: 无 (表头底线用cell-level)
    insideH = OxmlElement('w:insideH')
    insideH.set(qn('w:val'), 'none')
    insideH.set(qn('w:sz'), '0')
    insideH.set(qn('w:space'), '0')
    insideH.set(qn('w:color'), 'auto')
    borders.append(insideH)

    tblPr.append(borders)

    # === 表头行底部细线 0.75pt (cell级别) ===
    if len(rows_data) > 0:
        for j in range(num_cols):
            cell = table.rows[0].cells[j]
            tcPr = cell._tc.get_or_add_tcPr()
            # 移除旧tcBorders
            for old_tc in tcPr.findall(qn('w:tcBorders')):
                tcPr.remove(old_tc)
            tcBorders = OxmlElement('w:tcBorders')
            bb = OxmlElement('w:bottom')
            bb.set(qn('w:val'), 'single')
            bb.set(qn('w:sz'), '6')  # 0.75pt
            bb.set(qn('w:space'), '0')
            bb.set(qn('w:color'), '000000')
            tcBorders.append(bb)
            tcPr.append(tcBorders)

    return table

# ======================== 主解析逻辑 ========================
with open(MD_FILE, "r", encoding="utf-8") as f:
    lines = f.readlines()

# 状态机
in_abstract   = False   # 摘要正文区域
in_references = False   # 参考文献区域
in_table      = False   # 表格解析中
table_rows    = []
in_code       = False
skip_sep      = False
first_chapter = True    # 第一个章标题用分节符

for line in lines:
    stripped = line.strip()

    # --- 代码块 ---
    if stripped.startswith('```'):
        in_code = not in_code
        continue
    if in_code:
        continue

    # --- 水平线 ---
    if stripped in ('---', '---'):
        in_abstract = False  # 摘要区域结束
        continue

    # --- 图片 ---
    img_match = re.match(r'!\[(.*?)\]\((.+?)\)', stripped)
    if img_match:
        add_figure(img_match.group(2), img_match.group(1))
        continue

    # --- 图注行（已在add_figure中处理标题，此处跳过）---
    if re.match(r'^\*\*图\s+\d+-\d+\s+', stripped):
        continue

    # --- 表格分隔行 ---
    if re.match(r'^\|[-:\s|]+\|$', stripped):
        skip_sep = True
        continue

    # --- 表格行 ---
    if stripped.startswith('|') and stripped.endswith('|'):
        cells = [c.strip() for c in stripped.split('|')[1:-1]]
        if not in_table:
            in_table = True
            table_rows = [cells]
        else:
            if not skip_sep:
                table_rows.append(cells)
            skip_sep = False
        continue
    else:
        if in_table and table_rows:
            create_three_line_table(table_rows, len(table_rows[0]))
        in_table = False
        table_rows = []
        skip_sep = False

    # --- 空行 ---
    if not stripped:
        continue

    # --- 表注行 ---
    tc_match = re.match(r'^\*\*表\s+(\d+-\d+)\s+(.+?)\*\*$', stripped)
    if tc_match:
        add_table_caption(f"表 {tc_match.group(1)} {tc_match.group(2)}")
        continue

    # --- 标题 ---
    heading_match = re.match(r'^(#{1,4})\s+(.*)', stripped)
    if heading_match:
        level_md = len(heading_match.group(1))
        text = heading_match.group(2).strip()
        in_abstract = False  # 任何标题都结束摘要区域

        if text == '摘要':
            add_abstract_heading(text)
            in_abstract = True
            continue

        if text == '参考文献':
            in_references = True
            add_ref_heading(text)
            continue

        chapter_match = re.match(r'^第(\d+)章\s+', text)
        if chapter_match:
            in_references = False
            add_chapter_heading(text, is_first=first_chapter)
            first_chapter = False
        elif level_md == 1:
            add_chapter_heading(text, is_first=first_chapter)
            first_chapter = False
        elif level_md == 2:
            add_section_heading(text)
        elif level_md == 3:
            add_subsection_heading(text)
        elif level_md == 4:
            add_subsubsection_heading(text)
        continue

    # --- 关键词 ---
    kw_match = re.match(r'^\*\*关键词\*\*[：:]\s*(.*)', stripped)
    if kw_match:
        add_keywords_line(kw_match.group(1))
        in_abstract = False
        continue

    # --- 参考文献条目 ---
    if in_references and re.match(r'^\[(\d+)\]\s', stripped):
        add_reference_entry(stripped)
        continue

    # --- 跳过已经处理过的标注行 ---
    if re.match(r'^\*\*(关键词|图|表)', stripped):
        continue

    # --- 正文段落 ---
    if in_abstract:
        add_abstract_body(stripped)
    else:
        add_body_paragraph(stripped)

# 末尾残留表格
if in_table and table_rows:
    create_three_line_table(table_rows, len(table_rows[0]))

# ======================== 页码 ========================
# Section 0 (front matter: 摘要+关键词): 无页码
# Section 1+ (正文+参考文献): 页码 TNR 五号, 底部居中, 阿拉伯数字从1开始
for sec_idx, sec in enumerate(doc.sections):
    footer = sec.footer
    footer.is_linked_to_previous = False

    if sec_idx == 0:
        # 前导页（摘要）不编页码
        continue

    fp = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.first_line_indent = Cm(0)

    run = fp.add_run()
    set_font(run, cn='宋体', size=SZ_WUHAO)

    fldChar_begin = OxmlElement('w:fldChar')
    fldChar_begin.set(qn('w:fldCharType'), 'begin')
    run._r.append(fldChar_begin)

    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = ' PAGE '
    run._r.append(instrText)

    fldChar_end = OxmlElement('w:fldChar')
    fldChar_end.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar_end)

    # 设置页码从1开始
    sectPr = sec._sectPr
    pgNumType = OxmlElement('w:pgNumType')
    pgNumType.set(qn('w:start'), '1')
    sectPr.append(pgNumType)

# ======================== 保存 ========================
doc.save(DOCX_FILE)
print(f"[OK] DOCX saved: {DOCX_FILE}")
print(f"   File size: {os.path.getsize(DOCX_FILE) / 1024:.0f} KB")
print(f"   Format: 滇池学院期末考试论文格式规范 v2.1")
