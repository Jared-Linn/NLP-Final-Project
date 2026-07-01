"""
Complete reformatting of 学术论文.docx per 期末考试格式要求.txt
"""
from docx import Document
from docx.shared import Pt, Cm, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml
import re, os, copy

# ═══ Constants ═══
# Chinese font sizes: 三号=16pt, 四号=14pt, 小四=12pt, 五号=10.5pt
SIZE_TITLE   = Pt(22)  # paper title
SIZE_3       = Pt(16)  # 三号 (TOC title)
SIZE_4       = Pt(14)  # 四号 (chapter, references title)
SIZE_SMALL4  = Pt(12)  # 小四 (body, section headings, abstract title)
SIZE_5       = Pt(10.5) # 五号 (abstract body, keywords, table content, captions)

DOCX = r"E:\work\Claude code default\自然语言处理期末\论文\学术论文.docx"
OUT  = r"E:\work\Claude code default\自然语言处理期末\论文\学术论文_formatted.docx"

doc = Document(DOCX)

# ═══ Helper functions ═══
def cn_font_size(size):
    """Pt to EMU for comparison"""
    return size.emu if hasattr(size, 'emu') else int(size * 12700)

def set_run_font(run, cn='宋体', en='Times New Roman', size=SIZE_SMALL4, bold=False):
    run.font.size = size
    run.bold = bold
    run.font.name = en
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} />')
        rPr.insert(0, rFonts)
    rFonts.set(qn('w:eastAsia'), cn)
    rFonts.set(qn('w:ascii'), en)
    rFonts.set(qn('w:hAnsi'), en)

def set_para_font(para, cn='宋体', en='Times New Roman', size=SIZE_SMALL4, bold=False):
    for run in para.runs:
        set_run_font(run, cn, en, size, bold)

def set_para_spacing(para, before_pt=0, after_pt=0, line_spacing=1.5):
    pf = para.paragraph_format
    pf.space_before = Pt(before_pt)
    pf.space_after = Pt(after_pt)
    pf.line_spacing = line_spacing

def set_para_align(para, align):
    para.alignment = align

def clear_indent(para):
    para.paragraph_format.first_line_indent = None
    para.paragraph_format.left_indent = None

def set_first_indent(para, chars=2):
    """2-char first line indent at 12pt = 24pt"""
    para.paragraph_format.first_line_indent = Pt(chars * 12)

def set_hanging_indent(para, chars=2):
    """Hanging indent for references"""
    para.paragraph_format.left_indent = Pt(chars * 12)
    para.paragraph_format.first_line_indent = Pt(-chars * 12)

def add_page_break_before(para):
    pPr = para._element.get_or_add_pPr()
    for old in pPr.findall(qn('w:pageBreakBefore')):
        pPr.remove(old)
    pPr.append(parse_xml(f'<w:pageBreakBefore {nsdecls("w")} />'))

# ═══ Pattern matchers ═══
RE_CHAPTER   = re.compile(r'^第[一二三四五六七八九十\d]+\s*章\s')
RE_SECTION   = re.compile(r'^\d+\.\d+\s')       # X.X section
RE_SUBSECT   = re.compile(r'^\d+\.\d+\.\d+\s')   # X.X.X subsection
RE_FIG_CAP   = re.compile(r'^图\s*\d+[-]\d+')
RE_TBL_CAP   = re.compile(r'^表\s*\d+[-]\d+')
RE_REF_ITEM  = re.compile(r'^\[\d+\]')
RE_KEYWORD   = re.compile(r'^关键词')

# ═══ Phase 1: classify paragraphs ═══
classifications = []
found_abstract_title = False
found_keywords = False
in_abstract_body = False

for para in doc.paragraphs:
    t = para.text.strip()
    if not t:
        classifications.append('empty')
        continue

    # Detect structural position
    if t == '摘要':
        found_abstract_title = True
        classifications.append('abstract_title')
        continue

    if found_abstract_title and not found_keywords:
        if RE_KEYWORD.match(t):
            found_keywords = True
            classifications.append('keywords')
            continue
        else:
            classifications.append('abstract_body')
            continue

    # Classify by content pattern
    if t == '参考文献':
        classifications.append('references_title')
    elif RE_CHAPTER.match(t):
        classifications.append('chapter')
    elif RE_SUBSECT.match(t):
        classifications.append('subsection')
    elif RE_SECTION.match(t):
        classifications.append('section')
    elif RE_FIG_CAP.match(t):
        classifications.append('figure_caption')
    elif RE_TBL_CAP.match(t):
        classifications.append('table_caption')
    elif RE_REF_ITEM.match(t):
        classifications.append('reference_item')
    elif any(t.startswith(kw) for kw in ['摘要', '关键词', '目录']):
        classifications.append('meta')
    else:
        classifications.append('body')

print(f"Classified {len(classifications)} paragraphs:")
for cls in set(classifications):
    count = classifications.count(cls)
    print(f"  {cls}: {count}")

# ═══ Phase 2: Apply formatting ═══
for idx, (para, cls) in enumerate(zip(doc.paragraphs, classifications)):
    t = para.text.strip()

    if cls == 'empty':
        # Keep only strategic empty paragraphs
        # Remove empty body paragraphs (requirement: no blank lines between body text)
        # Find what comes before and after
        prev_cls = classifications[idx-1] if idx > 0 else None
        next_cls = classifications[idx+1] if idx+1 < len(classifications) else None

        # Keep spaces before/after: chapter headings, sections, figures, tables, keywords
        keep = False
        if next_cls in ('chapter', 'section', 'subsection', 'abstract_title', 'references_title'):
            keep = True
        if prev_cls in ('figure_caption', 'table_caption', 'abstract_title', 'keywords'):
            keep = True
        if next_cls in ('figure_caption', 'table_caption'):
            keep = True

        if not keep:
            # Remove the empty paragraph by clearing it (can't easily delete)
            # Set spacing to 0 and font to 1pt to minimize
            set_para_spacing(para, 0, 0, 1.0)
            if para.runs:
                set_run_font(para.runs[0], size=Pt(1))
            # Actually, let's try to remove the element
            para._element.getparent().remove(para._element)
        continue

    elif cls == 'abstract_title':
        clear_indent(para)
        set_para_align(para, WD_ALIGN_PARAGRAPH.CENTER)
        set_para_spacing(para, 12, 12, 1.0)  # 段前1行 段后1行 (at 12pt = 小四)
        set_para_font(para, '黑体', 'Times New Roman', SIZE_SMALL4, bold=True)

    elif cls == 'abstract_body':
        clear_indent(para)
        set_first_indent(para, 2)
        set_para_align(para, WD_ALIGN_PARAGRAPH.JUSTIFY)
        set_para_spacing(para, 0, 0, 1.5)
        set_para_font(para, '宋体', 'Times New Roman', SIZE_5, bold=False)

    elif cls == 'keywords':
        clear_indent(para)
        set_para_align(para, WD_ALIGN_PARAGRAPH.LEFT)
        set_para_spacing(para, 0, 0, 1.5)
        # "关键词" label: 黑体 小四 bold
        # Content: 宋体 五号 normal
        for run in para.runs:
            if '关键词' in run.text:
                set_run_font(run, '黑体', 'Times New Roman', SIZE_SMALL4, bold=True)
            else:
                set_run_font(run, '宋体', 'Times New Roman', SIZE_5, bold=False)

    elif cls == 'chapter':
        clear_indent(para)
        add_page_break_before(para)
        set_para_align(para, WD_ALIGN_PARAGRAPH.CENTER)
        set_para_spacing(para, 14, 7, 1.5)  # 段前1行 段后0.5行 (at 14pt)
        set_para_font(para, '黑体', 'Times New Roman', SIZE_4, bold=True)

    elif cls == 'section':
        clear_indent(para)
        set_para_align(para, WD_ALIGN_PARAGRAPH.LEFT)
        set_para_spacing(para, 6, 0, 1.5)  # 段前0.5行 (at 12pt)
        set_para_font(para, '黑体', 'Times New Roman', SIZE_SMALL4, bold=True)

    elif cls == 'subsection':
        clear_indent(para)
        set_para_align(para, WD_ALIGN_PARAGRAPH.LEFT)
        set_para_spacing(para, 6, 0, 1.5)  # 段前0.5行 (at 12pt)
        set_para_font(para, '宋体', 'Times New Roman', SIZE_SMALL4, bold=True)

    elif cls == 'figure_caption':
        clear_indent(para)
        set_para_align(para, WD_ALIGN_PARAGRAPH.CENTER)
        set_para_spacing(para, 0, 6, 1.5)  # 段后0.5行 (at 10.5pt)
        set_para_font(para, '黑体', 'Times New Roman', SIZE_5, bold=False)

    elif cls == 'table_caption':
        clear_indent(para)
        set_para_align(para, WD_ALIGN_PARAGRAPH.CENTER)
        set_para_spacing(para, 6, 0, 1.5)  # 段前0.5行 (at 10.5pt)
        set_para_font(para, '黑体', 'Times New Roman', SIZE_5, bold=False)

    elif cls == 'references_title':
        clear_indent(para)
        add_page_break_before(para)
        set_para_align(para, WD_ALIGN_PARAGRAPH.CENTER)
        set_para_spacing(para, 14, 7, 1.5)
        set_para_font(para, '黑体', 'Times New Roman', SIZE_4, bold=True)

    elif cls == 'reference_item':
        clear_indent(para)
        set_para_align(para, WD_ALIGN_PARAGRAPH.JUSTIFY)
        set_para_spacing(para, 0, 0, 1.5)
        set_para_font(para, '宋体', 'Times New Roman', SIZE_SMALL4, bold=False)
        set_hanging_indent(para, 2)

    elif cls == 'body':
        clear_indent(para)
        set_first_indent(para, 2)
        set_para_align(para, WD_ALIGN_PARAGRAPH.JUSTIFY)
        set_para_spacing(para, 0, 0, 1.5)
        set_para_font(para, '宋体', 'Times New Roman', SIZE_SMALL4, bold=False)

    elif cls == 'meta':
        pass  # skip

# ═══ Phase 3: Fix the paper title ═══
# First non-empty paragraph should be the title
for para in doc.paragraphs:
    t = para.text.strip()
    if t and t not in ('摘要', '目录'):
        clear_indent(para)
        set_para_align(para, WD_ALIGN_PARAGRAPH.CENTER)
        set_para_spacing(para, 24, 12, 1.5)
        set_para_font(para, '黑体', 'Times New Roman', SIZE_TITLE, bold=True)
        break

# ═══ Phase 4: Fix tables ═══
for table in doc.tables:
    # Triple-line table borders
    tbl = table._tbl
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = parse_xml(f'<w:tblPr {nsdecls("w")} />')
        tbl.insert(0, tblPr)

    # Set alignment
    jc = parse_xml(f'<w:jc {nsdecls("w")} w:val="center"/>')
    existing_jc = tblPr.findall(qn('w:jc'))
    for e in existing_jc:
        tblPr.remove(e)
    tblPr.append(jc)

    # Triple-line borders: thick top/bottom, thin under header, no sides/interior
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        '<w:top w:val="single" w:sz="12" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="single" w:sz="12" w:space="0" w:color="000000"/>'
        '<w:insideH w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        '<w:insideV w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        '<w:left w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        '<w:right w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        '</w:tblBorders>'
    )
    for old in tblPr.findall(qn('w:tblBorders')):
        tblPr.remove(old)
    tblPr.append(borders)

    # Style all cells
    for i, row in enumerate(table.rows):
        for cell in row.cells:
            for cell_para in cell.paragraphs:
                set_para_align(cell_para, WD_ALIGN_PARAGRAPH.CENTER)
                set_para_spacing(cell_para, 0, 0, 1.25)
                set_para_font(cell_para, '宋体', 'Times New Roman', SIZE_5)

            # Header row: add thin bottom border
            if i == 0:
                tcPr = cell._tc.get_or_add_tcPr()
                tcBorders = parse_xml(
                    f'<w:tcBorders {nsdecls("w")}>'
                    '<w:bottom w:val="single" w:sz="6" w:space="0" w:color="000000"/>'
                    '</w:tcBorders>'
                )
                for old in tcPr.findall(qn('w:tcBorders')):
                    tcPr.remove(old)
                tcPr.append(tcBorders)

# ═══ Phase 5: Center images ═══
for para in doc.paragraphs:
    drawings = para._element.findall('.//' + qn('w:drawing'))
    if drawings:
        clear_indent(para)
        set_para_align(para, WD_ALIGN_PARAGRAPH.CENTER)
        set_para_spacing(para, 6, 0, 1.5)

# ═══ Phase 6: Page numbers (Arabic, bottom center, 五号) ═══
for section in doc.sections:
    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    # Clear existing
    for r in fp.runs:
        r._element.getparent().remove(r._element)
    fp.clear()

    set_para_align(fp, WD_ALIGN_PARAGRAPH.CENTER)
    run = fp.add_run()
    set_run_font(run, 'Times New Roman', 'Times New Roman', SIZE_5)

    # PAGE field
    fld_begin = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>')
    fld_instr = parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve"> PAGE </w:instrText>')
    fld_end   = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>')
    run._element.append(fld_begin)
    run._element.append(fld_instr)
    run._element.append(fld_end)

# ═══ Save ═══
doc.save(OUT)
size_mb = os.path.getsize(OUT) / (1024*1024)
print(f"\nSaved: {OUT} ({size_mb:.1f} MB)")
